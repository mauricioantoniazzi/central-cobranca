"""Recontagem INDEPENDENTE (usa openpyxl, só para conferir; a Central em si não precisa de pacote nenhum): lê a planilha com openpyxl e refaz as contas do zero, sem importar nada da Central.
Definições tiradas do que as telas dizem, não do código delas. Valores em reais inteiros (a base só tem reais cheios)."""
import datetime as dt
import os
import math
import statistics as st
import warnings
from collections import defaultdict

import openpyxl

warnings.filterwarnings("ignore")
ARQ = os.environ.get("CENTRAL_BASE") or os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "Base_bruta.xlsx")
MESES = "janeiro fevereiro marco abril maio junho julho agosto setembro outubro novembro dezembro".split()


def _data_leia_me(wb):
    """Data de referência, lida por código PRÓPRIO daqui (não do da Central): '12 de agosto de 2026', 'dd/mm/aaaa' ou 'aaaa-mm-dd'."""
    import re
    import unicodedata
    for row in wb["Leia-me"].iter_rows(values_only=True):
        cels = [c for c in row if c is not None]
        if cels and "REFER" in str(cels[0]).upper() and len(cels) > 1:
            v = cels[1]
            if isinstance(v, (dt.date, dt.datetime)):
                return v.date() if isinstance(v, dt.datetime) else v
            t = "".join(c for c in unicodedata.normalize("NFD", str(v)) if unicodedata.category(c) != "Mn").lower()
            m = re.match(r"\s*(\d{1,2}) de (\w+)\.? de (\d{4})", t)
            if m:
                nome = m.group(2)
                n = next(i + 1 for i, mes in enumerate(MESES) if mes.startswith(nome[:3]))
                return dt.date(int(m.group(3)), n, int(m.group(1)))
            m = re.match(r"\s*(\d{1,2})/(\d{1,2})/(\d{4})", t)
            if m:
                return dt.date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
            m = re.match(r"\s*(\d{4})-(\d{2})-(\d{2})", t)
            if m:
                return dt.date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    raise SystemExit("recontagem independente: não achei a data de referência na aba Leia-me")


REF = dt.date(1970, 1, 1)  # preenchida por carregar() com a data da própria planilha


def carregar():
    wb = openpyxl.load_workbook(ARQ, data_only=True)
    global REF
    REF = _data_leia_me(wb)
    cli = {}
    for r in list(wb["Clientes"].iter_rows(values_only=True))[1:]:
        if r[0]:
            cli[int(r[0])] = {"nome": r[1], "prazo": int(r[6]), "limite": int(r[7]), "desde": r[8].date()}
    tit = []
    for r in list(wb["Títulos"].iter_rows(values_only=True))[1:]:
        if r[0]:
            tit.append({"num": r[0], "cli": int(r[1]), "emi": r[3].date(), "valor": int(r[4]), "venc": r[5].date(),
                        "pago": r[6].date() if r[6] else None})
    cob = []
    for r in list(wb["Cobranças"].iter_rows(values_only=True))[1:]:
        if r[0]:
            cob.append({"cli": int(r[1]), "quando": r[3], "valor": int(r[6])})
    return cli, tit, cob


def quantil(xs, p):
    xs = sorted(xs)
    k = (len(xs) - 1) * p
    f, c = math.floor(k), math.ceil(k)
    return xs[f] if f == c else xs[f] + (xs[c] - xs[f]) * (k - f)


def arred(x):
    return int(math.floor(x + 0.5))


def tudo():
    cli, tit, cob = carregar()
    R = {"ref": REF, "n_tit": len(tit), "n_cli": len(cli)}
    abertos = [t for t in tit if t["pago"] is None]
    venc = [t for t in abertos if t["venc"] < REF]
    R["aberto"] = sum(t["valor"] for t in abertos)
    R["n_aberto"] = len(abertos)
    R["vencido"] = sum(t["valor"] for t in venc)
    R["n_vencido"] = len(venc)
    R["n_cli_vencido"] = len({t["cli"] for t in venc})
    R["a_vencer"] = sum(t["valor"] for t in abertos if t["venc"] >= REF)
    ini = REF - dt.timedelta(days=6)
    R["entrou_7d"] = sum(t["valor"] for t in tit if t["pago"] and ini <= t["pago"] <= REF)
    R["n_entrou_7d"] = sum(1 for t in tit if t["pago"] and ini <= t["pago"] <= REF)
    R["n_cli_entrou_7d"] = len({t["cli"] for t in tit if t["pago"] and ini <= t["pago"] <= REF})
    ini8 = ini - dt.timedelta(days=56)
    R["media_8s"] = sum(t["valor"] for t in tit if t["pago"] and ini8 <= t["pago"] < ini) / 8
    R["venceu_7d"] = sum(t["valor"] for t in tit if ini <= t["venc"] <= REF)
    R["n_venceu_7d"] = sum(1 for t in tit if ini <= t["venc"] <= REF)
    R["venceu_7d_aberto"] = sum(t["valor"] for t in tit if ini <= t["venc"] <= REF and t["pago"] is None)
    R["n_venceu_7d_aberto"] = sum(1 for t in tit if ini <= t["venc"] <= REF and t["pago"] is None)

    fat12 = {}
    um_ano = REF - dt.timedelta(days=365)
    for c in cli:
        fat12[c] = sum(t["valor"] for t in tit if t["cli"] == c and um_ano < t["emi"] <= REF)
    total_fat = sum(fat12.values())

    C = {}
    for c, d in cli.items():
        ts = [t for t in tit if t["cli"] == c]
        if not ts:  # cliente sem nenhum título não tem o que recontar (a Central também o deixa de fora e avisa)
            continue
        pagos = [t for t in ts if t["pago"]]
        atr = [(t["pago"] - t["venc"]).days for t in pagos]
        ab = [t for t in ts if t["pago"] is None]
        vc = [t for t in ab if t["venc"] < REF]
        x = {"nome": d["nome"], "limite": d["limite"], "prazo": d["prazo"], "n_tit": len(ts), "n_pagos": len(pagos),
             "saldo": sum(t["valor"] for t in ab), "n_aberto": len(ab), "vencido": sum(t["valor"] for t in vc),
             "n_vencido": len(vc), "mais_antigo": max([(REF - t["venc"]).days for t in vc], default=0),
             "ocupacao": sum(t["valor"] for t in ab) / d["limite"], "fat12": fat12[c], "pct_fat": fat12[c] / total_fat,
             "primeira_compra": min(t["emi"] for t in ts), "ultima_compra": max(t["emi"] for t in ts),
             "dias_sem_comprar": (REF - max(t["emi"] for t in ts)).days, "desde": d["desde"]}
        x["pct_vencido_limite"] = x["vencido"] / d["limite"]
        if atr:
            x.update({"med_todos": st.median(atr), "p10_todos": quantil(atr, .1), "p90_todos": quantil(atr, .9), "max": max(atr),
                      "min": min(atr), "pct_atrasados": sum(1 for a in atr if a > 0) / len(atr)})
        # sazonalidade: mês do calendário (do vencimento) que dispara nos dois anos
        por = defaultdict(list)
        for t in pagos:
            por[(t["venc"].year, t["venc"].month)].append((t["pago"] - t["venc"]).days)
        if len(atr) >= 20:
            med = st.median(atr)
            mad = max(1.4826 * st.median(abs(a - med) for a in atr), 2.0)
            anos = defaultdict(set)
            for (a_, m_), xs in por.items():
                if st.median(xs) - med >= max(7, 1.5 * mad):
                    anos[m_].add(a_)
            saz = sorted(m for m, a in anos.items() if len(a) >= 2)
            x["sazonais"] = saz
            ns = [(t["pago"] - t["venc"]).days for t in pagos if t["venc"].month not in saz]
            x.update({"med_ns": st.median(ns), "p10_ns": quantil(ns, .1), "p90_ns": quantil(ns, .9), "n_ns": len(ns)})
            corte = REF - dt.timedelta(days=90)
            base = [(t["pago"] - t["venc"]).days for t in pagos if t["venc"].month not in saz and t["venc"] < corte]
            recp = [(t["pago"] - t["venc"]).days for t in pagos if t["venc"].month not in saz and t["venc"] >= corte]
            reca = [(REF - t["venc"]).days for t in vc if t["venc"] >= corte and t["venc"].month not in saz]
            x["n_base"], x["n_rec_pagos"], x["n_rec_abertos"] = len(base), len(recp), len(reca)
            if len(base) >= 20:
                q1, q3 = quantil(base, .25), quantil(base, .75)
                sig = max((q3 - q1) / 1.349, 2.0)
                x["base_med"], x["base_sig"], x["base_p10"], x["base_p90"] = st.median(base), sig, quantil(base, .1), quantil(base, .9)
                if len(recp) + len(reca) >= 5:
                    d_ = st.median(recp + reca) - x["base_med"]
                    x["piora"] = d_ >= 7 and d_ / sig >= 2
                    x["rec_med_pagos"] = st.median(recp) if recp else None
                    x["rec_med_misto"] = st.median(recp + reca)
        # compras: janelas de 90 dias ancoradas na data da base
        def janela(i):
            fim, ini_ = REF - dt.timedelta(days=90 * i), REF - dt.timedelta(days=90 * (i + 1))
            return sum(t["valor"] for t in ts if ini_ < t["emi"] <= fim)
        njan = (REF - x["primeira_compra"]).days // 90 - 1
        if njan >= 4:
            bs = [janela(i) for i in range(1, njan + 1)]
            x["compras_rec"], x["compras_media"], x["compras_dp"], x["njan"] = janela(0), st.mean(bs), st.pstdev(bs), njan
            x["queda"] = 1 - janela(0) / st.mean(bs)
        datas = sorted({t["emi"] for t in ts})
        gaps = [(b - a).days for a, b in zip(datas, datas[1:])]
        if len(datas) >= 10:
            x["gap_max"], x["gap_med"] = max(gaps), st.median(gaps)
        # fins de mês em que teria aparecido na lista de vencidos
        fins, d0 = [], x["primeira_compra"].replace(day=1)
        while True:
            d0 = (d0.replace(day=28) + dt.timedelta(days=4)).replace(day=1)
            fim = d0 - dt.timedelta(days=1)
            if fim >= REF:
                break
            fins.append(fim)
        pv = min(t["venc"] for t in ts)
        fins = [f for f in fins if f >= pv]
        x["fins"] = len(fins)
        x["fins_vencido"] = sum(1 for f in fins if any(t["venc"] <= f and t["emi"] <= f and (t["pago"] is None or t["pago"] > f) for t in ts))
        C[c] = x
    R["C"] = C

    # previsão: planilha (soma de vencimentos), semana a semana
    fora = 0
    sem = [0, 0, 0, 0]
    for t in abertos:
        k = (t["venc"] - REF).days
        if k <= 0:
            sem[0] += t["valor"]
        elif k <= 28:
            sem[(k - 1) // 7] += t["valor"]
        else:
            fora += t["valor"]
    R["planilha_sem"], R["planilha_total"], R["fora_janela"] = sem, sum(sem), fora
    R["cobrancas"] = cob
    R["cli"], R["tit"] = cli, tit
    return R


if __name__ == "__main__":
    R = tudo()
    for k in ("aberto", "n_aberto", "vencido", "n_vencido", "n_cli_vencido", "a_vencer", "entrou_7d", "n_entrou_7d", "n_cli_entrou_7d",
              "media_8s", "venceu_7d", "n_venceu_7d", "venceu_7d_aberto", "n_venceu_7d_aberto", "planilha_sem", "planilha_total", "fora_janela"):
        print(k, R[k])
    for c, x in R["C"].items():
        print(c, x["nome"][:22], {k: (round(v, 3) if isinstance(v, float) else v) for k, v in x.items()
                                   if k in ("saldo", "vencido", "n_vencido", "mais_antigo", "med_todos", "med_ns", "p10_ns", "p90_ns", "max", "sazonais",
                                            "base_med", "rec_med_pagos", "rec_med_misto", "piora", "queda", "dias_sem_comprar", "gap_max", "fins", "fins_vencido")})
